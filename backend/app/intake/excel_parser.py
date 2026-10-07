"""Spreadsheet parsing for packaging intake attachments."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from backend.app.intake.models import AttachmentExtraction, PackagingLineItem


def _normalize_header(value: Any, fallback: str) -> str:
    raw = str(value or fallback).strip().lower()
    return raw.replace(" ", "_").replace("-", "_")


def parse_spreadsheet_bytes(
    *,
    file_name: str,
    content: bytes,
    content_type: str | None = None,
) -> AttachmentExtraction:
    """Parse CSV/XLSX-like bytes into simple row dictionaries.

    XLSX parsing is optional and requires openpyxl. The app can still run without
    it, returning a clear attachment extraction error for the user/demo.
    """
    suffix = Path(file_name).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return AttachmentExtraction(
            file_name=file_name,
            content_type=content_type,
            rows=[dict(row) for row in reader],
        )

    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return AttachmentExtraction(
                file_name=file_name,
                content_type=content_type,
                extraction_status="error",
                errors=["openpyxl is required to parse Excel files."],
            )

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers_raw = next(iterator, None)
            if not headers_raw:
                continue
            headers = [
                _normalize_header(header, f"column_{index + 1}")
                for index, header in enumerate(headers_raw)
            ]
            for values in iterator:
                row = {
                    headers[index]: value
                    for index, value in enumerate(values)
                    if index < len(headers) and value is not None
                }
                if row:
                    row["_sheet"] = sheet.title
                    rows.append(row)

        return AttachmentExtraction(file_name=file_name, content_type=content_type, rows=rows)

    return AttachmentExtraction(
        file_name=file_name,
        content_type=content_type,
        extraction_status="unsupported",
        errors=[f"Unsupported file type '{suffix or 'unknown'}'. Upload CSV or XLSX."],
    )


def rows_to_line_items(rows: list[dict[str, Any]]) -> list[PackagingLineItem]:
    """Map spreadsheet rows to packaging line items using tolerant header aliases."""
    items: list[PackagingLineItem] = []
    for index, row in enumerate(rows, start=1):
        description = _first_present(row, ["description", "item", "item_name", "product", "material"])
        if not description:
            continue

        quantity_raw = _first_present(row, ["quantity", "qty", "required_quantity", "volume"])
        quantity = _to_float(quantity_raw)
        target_price = _to_float(_first_present(row, ["target_price", "budget", "unit_price"]))

        items.append(
            PackagingLineItem(
                item_number=index,
                description=str(description),
                packaging_category=_first_present(row, ["packaging_category", "category", "type"]),
                quantity=quantity,
                unit=_first_present(row, ["unit", "uom"]),
                material=_first_present(row, ["material", "paper_type"]),
                gsm=_string_or_none(_first_present(row, ["gsm"])),
                bf=_string_or_none(_first_present(row, ["bf", "bursting_factor"])),
                ply=_string_or_none(_first_present(row, ["ply", "no_of_ply"])),
                flute_type=_string_or_none(_first_present(row, ["flute", "flute_type"])),
                color=_string_or_none(_first_present(row, ["color", "colour"])),
                printing_details=_string_or_none(_first_present(row, ["printing", "print", "artwork"])),
                usage_context=_string_or_none(_first_present(row, ["usage", "use_case", "usage_context"])),
                specifications=_string_or_none(_first_present(row, ["specifications", "specs", "notes"])),
                target_price=target_price,
            )
        )
    return items


def _first_present(row: dict[str, Any], keys: list[str]) -> Any:
    normalized = {_normalize_header(key, key): value for key, value in row.items()}
    for key in keys:
        value = normalized.get(key)
        if value not in (None, ""):
            return value
    return None


def _to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def _string_or_none(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return str(value)
